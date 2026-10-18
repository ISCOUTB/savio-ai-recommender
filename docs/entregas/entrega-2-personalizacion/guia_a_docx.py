"""Genera el entregable .docx de la guia paso a paso (E2-02) a partir del Markdown.

Uso (desde la raiz del repo):
    python docs/entregas/entrega-2-personalizacion/guia_a_docx.py

Requiere: python-docx (pip install python-docx).
"""

import re
import sys
from pathlib import Path

from docx import Document

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(__file__).resolve().parent
SRC = BASE / "guia-paso-a-paso-docker-moodle.md"
DST = BASE / "Guia_Paso_a_Paso_Docker_Moodle_E2_02.docx"

_CODE_FONT = "Consolas"


def _style_exists(doc, name):
    try:
        doc.styles[name]
        return True
    except KeyError:
        return False


def add_runs(paragraph, text):
    """Agrega texto con formato inline basico: **negrita** y `codigo`."""
    for token in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = _CODE_FONT
        else:
            paragraph.add_run(token)
    return paragraph


def add_code_block(doc, code_lines):
    paragraph = doc.add_paragraph()
    if _style_exists(doc, "No Spacing"):
        paragraph.style = doc.styles["No Spacing"]
    run = paragraph.add_run("\n".join(code_lines))
    run.font.name = _CODE_FONT
    return paragraph


def add_table(doc, raw_rows):
    body = [
        cells
        for cells in raw_rows
        if not all(re.match(r"^:?-+:?$", c) for c in cells)
    ]
    if not body:
        return
    ncols = max(len(cells) for cells in body)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    for row_index, cells in enumerate(body):
        row = table.add_row().cells
        for col_index in range(ncols):
            text = cells[col_index] if col_index < len(cells) else ""
            row[col_index].text = ""
            add_runs(row[col_index].paragraphs[0], text)
            if row_index == 0:
                for run in row[col_index].paragraphs[0].runs:
                    run.bold = True


def convert(md_path, docx_path):
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    index, total = 0, len(lines)
    in_code, code_buffer = False, []

    while index < total:
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            if not in_code:
                in_code, code_buffer = True, []
            else:
                in_code = False
                add_code_block(doc, code_buffer)
            index += 1
            continue

        if in_code:
            code_buffer.append(line)
            index += 1
            continue

        if not stripped or stripped == "---":
            index += 1
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            doc.add_heading(heading.group(2), level=min(len(heading.group(1)), 4))
            index += 1
            continue

        if stripped.startswith("|"):
            raw_rows = []
            while index < total and lines[index].strip().startswith("|"):
                cells = [c.strip() for c in lines[index].strip().strip("|").split("|")]
                raw_rows.append(cells)
                index += 1
            add_table(doc, raw_rows)
            continue

        if stripped.startswith("> "):
            add_runs(doc.add_paragraph(style="Intense Quote"), stripped[2:])
            index += 1
            continue

        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        if bullet:
            add_runs(doc.add_paragraph(style="List Bullet"), bullet.group(1))
            index += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if numbered:
            add_runs(doc.add_paragraph(style="List Number"), numbered.group(1))
            index += 1
            continue

        add_runs(doc.add_paragraph(), stripped)
        index += 1

    doc.save(docx_path)
    return docx_path


if __name__ == "__main__":
    if not SRC.exists():
        raise SystemExit(f"No se encontro la fuente: {SRC}")
    out = convert(SRC, DST)
    print(f"OK -> {out}")
